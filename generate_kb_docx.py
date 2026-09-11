"""
Convert career_companion_knowledge.md to a styled docx file:
knowledge_base/career_companion_knowledge.docx
"""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def md_to_docx():
    root = Path(__file__).resolve().parent
    md_path = root / "knowledge_base" / "career_companion_knowledge.md"
    docx_path = root / "knowledge_base" / "career_companion_knowledge.docx"

    if not md_path.exists():
        print(f"Error: {md_path} not found")
        return

    doc = Document()

    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    for line in lines:
        raw = line.rstrip()

        # Empty line
        if not raw:
            continue

        # H1
        if raw.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(raw[2:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x11, 0x33, 0x66)
            continue

        # H2
        if raw.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(raw[3:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            continue

        # H3
        if raw.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(raw[4:].strip())
            run.font.name = 'Calibri'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
            continue

        # Horizontal rule or quote block
        if raw.startswith("---") or raw.startswith(">"):
            if raw.startswith(">"):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(raw[1:].strip())
                run.font.italic = True
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            continue

        # Bullet lists (- or *)
        if re.match(r"^[-*]\s+", raw):
            text = re.sub(r"^[-*]\s+", "", raw).strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            _add_formatted_text(p, text)
            continue

        # Numbered list
        if re.match(r"^\d+\.\s+", raw):
            text = re.sub(r"^\d+\.\s+", "", raw).strip()
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            _add_formatted_text(p, text)
            continue

        # Normal paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        _add_formatted_text(p, raw)

    doc.save(str(docx_path))
    print(f"Successfully generated docx at: {docx_path}")

def _add_formatted_text(paragraph, text: str):
    # Process simple bold **text**
    tokens = re.split(r"(\*\*.*?\*\*)", text)
    for token in tokens:
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.font.bold = True
        else:
            paragraph.add_run(token)

if __name__ == "__main__":
    md_to_docx()
