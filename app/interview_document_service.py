"""
Interview Document Processing & Retrieval Service (Phase B).

Handles:
1. Validated text extraction from PDF (pdfplumber) and DOCX (python-docx with tables).
2. Detection of empty / scanned / unextractable documents.
3. Semantic text chunking and isolated vector indexing per user.
4. User-isolated semantic retrieval and representative overview chunking.
5. Active document lifecycle (upload, query, deactivate/remove).
"""

import json
import os
from typing import Any, Dict, List, Optional, Tuple

import numpy as np
from sqlalchemy.orm import Session

from app import models
from app.embeddings import model as embedding_model


UPLOAD_DOC_DIR = os.path.join("uploads", "interview_documents")
os.makedirs(UPLOAD_DOC_DIR, exist_ok=True)

CHUNK_SIZE = 700
CHUNK_OVERLAP = 120
MAX_SUMMARY_CHUNKS = 5


# ──────────────────────────────────────────────────────────────────
# 1. Text Extraction
# ──────────────────────────────────────────────────────────────────

def extract_text_from_pdf(file_path: str) -> str:
    """
    Extract text page by page from a PDF using pdfplumber.
    Raises ValueError if PDF is scanned, empty, or unextractable.
    """
    import pdfplumber

    text_parts = []
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text and page_text.strip():
                    text_parts.append(page_text.strip())
    except Exception as e:
        raise ValueError(f"Failed to read PDF file: {str(e)}")

    full_text = "\n\n".join(text_parts).strip()
    if not full_text:
        raise ValueError("This PDF does not contain extractable text. Please upload a text-based PDF.")

    return full_text


def extract_text_from_docx(file_path: str) -> str:
    """
    Extract text from a DOCX document including paragraph text and table cell text.
    Raises ValueError if document is empty.
    """
    from docx import Document

    text_parts = []
    try:
        doc = Document(file_path)

        # Extract paragraphs
        for p in doc.paragraphs:
            val = p.text.strip()
            if val:
                text_parts.append(val)

        # Extract tables
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    text_parts.append(" | ".join(cells))

    except Exception as e:
        raise ValueError(f"Failed to read DOCX file: {str(e)}")

    full_text = "\n\n".join(text_parts).strip()
    if not full_text:
        raise ValueError("The uploaded DOCX document contains no extractable text.")

    return full_text


def extract_document_text(file_path: str, extension: str) -> str:
    """Route document extraction by extension."""
    ext = extension.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported document format: {extension}. Only .pdf and .docx are supported.")


# ──────────────────────────────────────────────────────────────────
# 2. Chunking & Embeddings
# ──────────────────────────────────────────────────────────────────

def chunk_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> List[str]:
    """
    Split text into overlapping semantic chunks.
    Respects paragraph boundaries where possible.
    """
    if not text:
        return []

    # If small enough, single chunk is sufficient
    if len(text) <= chunk_size:
        return [text.strip()]

    paragraphs = text.split("\n\n")
    chunks: List[str] = []
    current_chunk = ""

    for p in paragraphs:
        p = p.strip()
        if not p:
            continue

        if len(current_chunk) + len(p) + 2 <= chunk_size:
            current_chunk = f"{current_chunk}\n\n{p}" if current_chunk else p
        else:
            if current_chunk:
                chunks.append(current_chunk.strip())
                # Start next chunk with overlap from the tail of current_chunk
                if len(current_chunk) > overlap:
                    current_chunk = current_chunk[-overlap:] + "\n\n" + p
                else:
                    current_chunk = p
            else:
                # A single huge paragraph: break by character length
                start = 0
                while start < len(p):
                    end = start + chunk_size
                    chunks.append(p[start:end].strip())
                    start = end - overlap

    if current_chunk and current_chunk.strip():
        chunks.append(current_chunk.strip())

    return chunks if chunks else [text[:chunk_size]]


def save_and_index_document(
    user_id: int,
    original_filename: str,
    file_path: str,
    extension: str,
    extracted_text: str,
    db: Session,
) -> models.InterviewDocument:
    """
    Deactivate previous active documents for this user, create new InterviewDocument,
    chunk text, compute embeddings, and store chunk records.
    Strictly isolated to user_id.
    """
    # 1. Mark existing active documents for this user as inactive
    db.query(models.InterviewDocument).filter(
        models.InterviewDocument.user_id == user_id,
        models.InterviewDocument.is_active == True,
    ).update({"is_active": False})

    # 2. Chunk text
    chunks = chunk_text(extracted_text)

    # 3. Create document record
    doc = models.InterviewDocument(
        user_id=user_id,
        filename=original_filename,
        file_path=file_path,
        file_type=extension.lstrip(".").lower(),
        extracted_text=extracted_text,
        chunk_count=len(chunks),
        is_active=True,
    )
    db.add(doc)
    db.flush()  # assign doc.id

    # 4. Generate batch embeddings
    if chunks:
        embeddings = embedding_model.encode(chunks, show_progress_bar=False).tolist()

        for idx, (c_text, emb) in enumerate(zip(chunks, embeddings)):
            chunk_rec = models.InterviewDocumentChunk(
                document_id=doc.id,
                user_id=user_id,
                chunk_index=idx,
                chunk_text=c_text,
                embedding=json.dumps(emb),
            )
            db.add(chunk_rec)

    db.commit()
    db.refresh(doc)
    return doc


# ──────────────────────────────────────────────────────────────────
# 3. Isolated Document Retrieval
# ──────────────────────────────────────────────────────────────────

def get_active_document(user_id: int, db: Session) -> Optional[models.InterviewDocument]:
    """Retrieve the currently active document for the authenticated user."""
    return (
        db.query(models.InterviewDocument)
        .filter(
            models.InterviewDocument.user_id == user_id,
            models.InterviewDocument.is_active == True,
        )
        .order_by(models.InterviewDocument.uploaded_at.desc())
        .first()
    )


def retrieve_document_context(
    user_id: int,
    user_message: str,
    db: Session,
    top_k: int = 5,
) -> Tuple[Optional[models.InterviewDocument], List[str], str]:
    """
    Retrieve relevant context from the authenticated user's active document.
    Returns (document, list_of_chunks, formatted_context_string).
    Guarantees strict user data isolation.
    """
    doc = get_active_document(user_id, db)
    if not doc:
        return None, [], ""

    # If small document (<= 2500 characters), pass full text directly
    if len(doc.extracted_text) <= 2500:
        formatted = (
            f"=== UPLOADED DOCUMENT: '{doc.filename}' ({doc.file_type.upper()}) ===\n"
            f"{doc.extracted_text}\n"
            f"============================================================"
        )
        return doc, [doc.extracted_text], formatted

    # For larger documents, retrieve chunks
    chunks_records = (
        db.query(models.InterviewDocumentChunk)
        .filter(
            models.InterviewDocumentChunk.document_id == doc.id,
            models.InterviewDocumentChunk.user_id == user_id,
        )
        .order_by(models.InterviewDocumentChunk.chunk_index.asc())
        .all()
    )

    if not chunks_records:
        # Fallback to initial snippet
        snippet = doc.extracted_text[:2000]
        formatted = (
            f"=== UPLOADED DOCUMENT: '{doc.filename}' ({doc.file_type.upper()}) ===\n"
            f"{snippet}\n"
            f"============================================================"
        )
        return doc, [snippet], formatted

    msg_lower = user_message.lower()
    is_overview_query = any(
        k in msg_lower for k in [
            "summarize", "summary", "overview", "roadmap", "study plan",
            "preparation plan", "important topics", "what topics", "outline",
            "what is this document about", "give me questions and answers",
            "generate questions and answers", "generate interview questions",
        ]
    )

    selected_chunks: List[str] = []

    if is_overview_query:
        # Select representative chunks across the document (beginning, evenly distributed)
        total = len(chunks_records)
        if total <= MAX_SUMMARY_CHUNKS:
            selected_chunks = [c.chunk_text for c in chunks_records]
        else:
            indices = np.linspace(0, total - 1, MAX_SUMMARY_CHUNKS, dtype=int).tolist()
            selected_chunks = [chunks_records[i].chunk_text for i in indices]
    else:
        # Targeted question: Semantic similarity search with numpy
        try:
            q_emb = np.array(embedding_model.encode([user_message], show_progress_bar=False)[0], dtype=np.float32)
            q_norm = np.linalg.norm(q_emb)

            scored_chunks = []
            for c in chunks_records:
                if c.embedding:
                    c_emb = np.array(json.loads(c.embedding), dtype=np.float32)
                    c_norm = np.linalg.norm(c_emb)
                    if q_norm > 0 and c_norm > 0:
                        sim = float(np.dot(q_emb, c_emb) / (q_norm * c_norm))
                    else:
                        sim = 0.0
                    scored_chunks.append((sim, c.chunk_text))
                else:
                    scored_chunks.append((0.0, c.chunk_text))

            scored_chunks.sort(key=lambda x: x[0], reverse=True)
            selected_chunks = [text for _, text in scored_chunks[:top_k]]
        except Exception as e:
            print(f"[interview_document_service] Semantic retrieval error: {e}")
            selected_chunks = [c.chunk_text for c in chunks_records[:top_k]]

    joined_chunks = "\n\n---\n\n".join(selected_chunks)
    formatted = (
        f"=== UPLOADED DOCUMENT: '{doc.filename}' ({doc.file_type.upper()}) ===\n"
        f"Relevant Document Excerpts:\n\n"
        f"{joined_chunks}\n"
        f"============================================================"
    )
    return doc, selected_chunks, formatted


# ──────────────────────────────────────────────────────────────────
# 4. Document Lifecycle
# ──────────────────────────────────────────────────────────────────

def remove_document(
    user_id: int,
    document_id: Optional[int],
    db: Session,
) -> bool:
    """
    Remove the active document (or specified document_id) for the authenticated user.
    Deletes associated chunks, removes file from disk, and removes document record.
    Strictly isolated: queries user_id == user_id.
    """
    query = db.query(models.InterviewDocument).filter(
        models.InterviewDocument.user_id == user_id
    )
    if document_id is not None:
        query = query.filter(models.InterviewDocument.id == document_id)
    else:
        query = query.filter(models.InterviewDocument.is_active == True)

    doc = query.first()
    if not doc:
        return False

    # Delete chunks
    db.query(models.InterviewDocumentChunk).filter(
        models.InterviewDocumentChunk.document_id == doc.id,
        models.InterviewDocumentChunk.user_id == user_id,
    ).delete()

    # Clean up physical file if it exists
    if doc.file_path and os.path.exists(doc.file_path):
        try:
            os.remove(doc.file_path)
        except Exception:
            pass

    # Delete document record
    db.delete(doc)
    db.commit()
    return True
